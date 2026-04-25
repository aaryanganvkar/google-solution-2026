# model.py - Full implementation with S3 + HuggingFace LLM
import os
import json
import uuid
import tempfile
from datetime import datetime
from typing import List, Dict, Any, Optional
from dataclasses import dataclass
from enum import Enum
import re
import warnings
from dotenv import load_dotenv

warnings.filterwarnings('ignore')
load_dotenv()

# ── AWS S3 ──────────────────────────────────────────────────────────────────
AWS_S3_BUCKET = os.getenv('AWS_S3_BUCKET')
AWS_REGION    = os.getenv('AWS_REGION', 'us-east-1')

def _s3():
    import boto3
    return boto3.client(
        's3',
        aws_access_key_id=os.getenv('AWS_ACCESS_KEY_ID'),
        aws_secret_access_key=os.getenv('AWS_SECRET_ACCESS_KEY'),
        region_name=AWS_REGION,
    )

# Keep a module-level client for backwards compat (used by processing_api)
try:
    import boto3
    s3_client = _s3()
except Exception:
    s3_client = None

# ── HuggingFace LLM ─────────────────────────────────────────────────────────
HF_TOKEN   = os.getenv('HF_TOKEN') or os.getenv('HUGGING_FACE_TOKEN', '')
MODEL_NAME = "mistralai/Mistral-7B-Instruct-v0.2"   # free-tier serverless chat model

try:
    from huggingface_hub import InferenceClient
    _hf_client = InferenceClient(api_key=HF_TOKEN) if HF_TOKEN else None
except Exception:
    _hf_client = None

# ── Enums / dataclasses ─────────────────────────────────────────────────────
class DocumentType(Enum):
    INVOICE           = "invoice"
    TECHNICAL_DOC     = "technical_documentation"
    TIMELINE          = "project_timeline"
    SAFETY_REPORT     = "safety_report"
    COMPLIANCE_DOC    = "compliance_document"
    HR_DOCUMENT       = "hr_document"
    ENGINEERING_REPORT= "engineering_report"
    OPERATIONS_MANUAL = "operations_manual"
    PROCUREMENT_ORDER = "procurement_order"
    ADMINISTRATIVE    = "administrative"
    UNKNOWN           = "unknown"

class Department(Enum):
    ENGINEERING = "engineering"
    OPERATIONS  = "operations"
    PROCUREMENT = "procurement"
    HR          = "hr"
    SAFETY      = "safety"
    COMPLIANCE  = "compliance"
    ADMIN       = "admin"
    FINANCE     = "finance"
    MANAGEMENT  = "management"

@dataclass
class DocumentProcessingResult:
    file_path:          str
    original_filename:  str
    processed_filename: str
    document_type:      DocumentType
    department:         Department
    summary:            str
    key_points:         List[str]
    action_items:       List[str]
    deadline:           Optional[str]
    priority:           str
    metadata:           Dict[str, Any]
    raw_text:           str
    processed_date:     str
    s3_key:             Optional[str] = None
    s3_url:             Optional[str] = None

@dataclass
class CalendarEvent:
    title:           str
    description:     str
    date:            str
    department:      str
    priority:        str
    action_required: bool

# ── LLM helper ──────────────────────────────────────────────────────────────
def call_llm(prompt: str, system_message: str = None, max_tokens: int = 800) -> str:
    """Call HuggingFace Inference API; falls back to rule-based if unavailable."""
    if not _hf_client:
        return ""
    messages = []
    if system_message:
        messages.append({"role": "system", "content": system_message})
    messages.append({"role": "user", "content": prompt})
    try:
        completion = _hf_client.chat.completions.create(
            model=MODEL_NAME,
            messages=messages,
            max_tokens=max_tokens,
            temperature=0.3,
        )
        return completion.choices[0].message.content.strip()
    except Exception as e:
        print(f"LLM call failed ({e}); using rule-based fallback.")
        return ""

# ── Text extraction ─────────────────────────────────────────────────────────
def extract_text_from_file(file_path: str) -> str:
    """Extract plain text from PDF, DOCX, XLSX, or plain-text files."""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == '.pdf':
            import fitz  # PyMuPDF
            doc  = fitz.open(file_path)
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return text

        elif ext in ('.docx', '.doc'):
            from docx import Document
            doc  = Document(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return text

        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb   = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(rows)

        else:  # .txt and everything else
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()

    except Exception as e:
        print(f"Text extraction failed for {file_path}: {e}")
        return ""

# ── Classification ───────────────────────────────────────────────────────────
_TYPE_KEYWORDS: Dict[DocumentType, List[str]] = {
    DocumentType.INVOICE:            ["invoice", "bill", "payment due", "total amount", "vendor", "po number"],
    DocumentType.SAFETY_REPORT:      ["safety", "incident", "hazard", "ppe", "accident", "risk assessment", "osha"],
    DocumentType.COMPLIANCE_DOC:     ["compliance", "regulation", "audit", "gdpr", "iso", "certification", "regulatory"],
    DocumentType.HR_DOCUMENT:        ["employee", "hire", "payroll", "leave", "performance review", "onboarding", "resignation"],
    DocumentType.ENGINEERING_REPORT: ["engineering", "design", "structural", "cad", "specification", "load", "blueprint"],
    DocumentType.OPERATIONS_MANUAL:  ["operations", "procedure", "sop", "maintenance", "workflow", "process"],
    DocumentType.PROCUREMENT_ORDER:  ["purchase order", "procurement", "supplier", "quotation", "rfq", "delivery"],
    DocumentType.TIMELINE:           ["timeline", "milestone", "gantt", "schedule", "project plan", "deadline"],
    DocumentType.ADMINISTRATIVE:     ["memo", "minutes", "agenda", "announcement", "policy", "directive"],
    DocumentType.TECHNICAL_DOC:      ["technical", "manual", "documentation", "api", "system", "software"],
}

def classify_document(text: str) -> DocumentType:
    lower = text.lower()

    # Try LLM first
    llm_resp = call_llm(
        f"Classify this document into EXACTLY one of these types: "
        f"invoice, safety_report, compliance_document, hr_document, engineering_report, "
        f"operations_manual, procurement_order, project_timeline, administrative, technical_documentation, unknown.\n\n"
        f"Document excerpt (first 1500 chars):\n{lower[:1500]}\n\n"
        f"Reply with only the type name, nothing else.",
        max_tokens=20,
    )
    for dt in DocumentType:
        if dt.value in llm_resp.lower():
            return dt

    # Rule-based fallback
    scores = {dt: 0 for dt in DocumentType}
    for dt, kws in _TYPE_KEYWORDS.items():
        for kw in kws:
            if kw in lower:
                scores[dt] += 1
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else DocumentType.UNKNOWN

def determine_department(doc_type: DocumentType, text: str) -> Department:
    mapping = {
        DocumentType.INVOICE:            Department.PROCUREMENT,
        DocumentType.SAFETY_REPORT:      Department.SAFETY,
        DocumentType.COMPLIANCE_DOC:     Department.COMPLIANCE,
        DocumentType.HR_DOCUMENT:        Department.HR,
        DocumentType.ENGINEERING_REPORT: Department.ENGINEERING,
        DocumentType.OPERATIONS_MANUAL:  Department.OPERATIONS,
        DocumentType.PROCUREMENT_ORDER:  Department.PROCUREMENT,
        DocumentType.TIMELINE:           Department.MANAGEMENT,
        DocumentType.ADMINISTRATIVE:     Department.ADMIN,
        DocumentType.TECHNICAL_DOC:      Department.ENGINEERING,
        DocumentType.UNKNOWN:            Department.ADMIN,
    }
    return mapping.get(doc_type, Department.ADMIN)

# ── Summary / key-points / action-items ─────────────────────────────────────
def create_summary(text: str, doc_type: DocumentType) -> str:
    excerpt = text[:3000]
    llm_resp = call_llm(
        f"Write a concise 2-3 sentence summary of this {doc_type.value} document:\n\n{excerpt}",
        system_message="You are a professional document analyst. Be factual and concise.",
        max_tokens=200,
    )
    if llm_resp:
        return llm_resp

    # Rule-based: skip header/metadata lines, find first real paragraph
    lines = text.splitlines()
    skip_patterns = re.compile(
        r'^(\s*[-•*]|\s*\d+[.)\s]|[A-Z\s]{5,}:|.{0,40}:\s|'  # bullet, number, ALL CAPS HEADER, key: value
        r'\[.\]|report id|date:|ref:|version:|approved|prepared by|signed)',
        re.IGNORECASE
    )
    meaningful = [
        l.strip() for l in lines
        if len(l.strip()) > 60
        and not skip_patterns.match(l.strip())
        and not l.strip().isupper()  # skip all-caps section titles
    ]
    # Take first 2 meaningful lines and join into a clean sentence
    summary_text = ' '.join(meaningful[:2])
    return summary_text if summary_text else "Document summary not available."

def extract_key_points(text: str) -> List[str]:
    llm_resp = call_llm(
        f"List up to 5 key points from this document as a bullet list (one per line, start each with '-'):\n\n{text[:3000]}",
        system_message="You are a professional document analyst. Be concise.",
        max_tokens=300,
    )
    if llm_resp:
        points = [
            line.lstrip("-•* ").strip()
            for line in llm_resp.splitlines()
            if line.strip() and len(line.strip()) > 10
        ]
        return points[:5] if points else _rule_based_key_points(text)
    return _rule_based_key_points(text)

def _rule_based_key_points(text: str) -> List[str]:
    """Extract structured key points — looks for bullet/numbered list items and section highlights."""
    points = []
    lines  = text.splitlines()

    for line in lines:
        stripped = line.strip()
        # Match bullet points or numbered list items with real content
        if re.match(r'^[-•*]\s+\S.{15,}', stripped):
            clean = re.sub(r'^[-•*]\s+', '', stripped)
            points.append(clean)
        elif re.match(r'^\d+[.)\s]\s*\S.{15,}', stripped):
            clean = re.sub(r'^\d+[.)\s]\s*', '', stripped)
            points.append(clean)
        # Short labelled facts: "Total: X", "Status: Y", "Priority: Z"
        elif re.match(r'^[A-Za-z ]{3,25}:\s+\S.{5,}', stripped) and len(stripped) < 120:
            points.append(stripped)

    # Deduplicate while preserving order
    seen   = set()
    unique = []
    for p in points:
        key = p.lower()[:40]
        if key not in seen:
            seen.add(key)
            unique.append(p)

    if unique:
        return unique[:5]
    # Last resort: first 4 long sentences that aren't document headers
    sentences = [
        s.strip() for s in re.split(r'(?<=[.!?])\s+', text)
        if 50 < len(s.strip()) < 200 and not s.strip().isupper()
    ]
    return sentences[:4] if sentences else ["Key points could not be extracted automatically."]

def extract_action_items(text: str) -> List[str]:
    llm_resp = call_llm(
        f"List all action items, tasks, or required actions from this document (one per line, start with '-'):\n\n{text[:3000]}",
        system_message="Extract only explicit action items or tasks. If none, reply 'No action items found.'",
        max_tokens=300,
    )
    if llm_resp and "no action" not in llm_resp.lower():
        items = [
            line.lstrip("-•* ").strip()
            for line in llm_resp.splitlines()
            if line.strip() and len(line.strip()) > 10
        ]
        return items[:5] if items else []

    # Rule-based: look for imperative sentences
    action_kws = ["must", "should", "required", "ensure", "complete", "submit", "review", "approve", "update"]
    sentences   = re.split(r'(?<=[.!?])\s+', text)
    items       = [s.strip() for s in sentences if any(kw in s.lower() for kw in action_kws)]
    return items[:4]

def extract_deadline(text: str) -> Optional[str]:
    """Return the single most relevant deadline date from the document."""
    # Patterns that are explicitly labelled as deadlines
    labelled_patterns = [
        r'(?:deadline|due date|must be (?:completed?|submitted?)|no later than|complete by)[:\s]+([A-Za-z]+ \d{1,2},? \d{4})',
        r'(?:deadline|due|by|before)[:\s]+(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
    ]
    for pat in labelled_patterns:
        match = re.search(pat, text, re.IGNORECASE)
        if match:
            return match.group(1).strip()

    # Fall back: any date that appears in the same sentence as action words
    action_words = re.compile(r'\b(submit|complete|review|approve|deliver|resolve|close|ensure|deadline)\b', re.I)
    date_pat     = re.compile(r'\b([A-Za-z]+ \d{1,2},? \d{4}|\d{1,2}[/-]\d{1,2}[/-]\d{4})\b')
    for sentence in re.split(r'(?<=[.!?])\s+', text):
        if action_words.search(sentence):
            date_match = date_pat.search(sentence)
            if date_match:
                return date_match.group(1).strip()

    return None

def determine_priority(text: str) -> str:
    lower = text.lower()
    high_kws   = ["urgent", "critical", "immediate", "asap", "emergency", "high priority", "mandatory"]
    medium_kws = ["important", "required", "necessary", "priority", "attention"]
    if any(kw in lower for kw in high_kws):
        return "high"
    if any(kw in lower for kw in medium_kws):
        return "medium"
    return "low"

# ── S3 helpers ───────────────────────────────────────────────────────────────
def download_from_s3(s3_key: str) -> str:
    client = _s3()
    temp   = tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(s3_key)[1])
    client.download_file(AWS_S3_BUCKET, s3_key, temp.name)
    print(f"✅ Downloaded from S3: {s3_key}")
    return temp.name

def upload_to_s3(file_path: str, department: str, document_type: str) -> Dict[str, str]:
    client    = _s3()
    filename  = os.path.basename(file_path)
    ts        = datetime.now().strftime("%Y%m%d_%H%M%S")
    uid       = uuid.uuid4().hex[:8]
    s3_key    = f"processed/{department}/{document_type}/{ts}_{uid}_{filename}"
    client.upload_file(
        file_path, AWS_S3_BUCKET, s3_key,
        ExtraArgs={
            'ContentType': 'application/octet-stream',
            'Metadata': {
                'department':    department,
                'document_type': document_type,
                'processed_date': datetime.now().isoformat(),
            }
        }
    )
    s3_url = f"https://{AWS_S3_BUCKET}.s3.{AWS_REGION}.amazonaws.com/{s3_key}"
    print(f"✅ Uploaded to S3: {s3_key}")
    return {'key': s3_key, 'url': s3_url}

def list_s3_documents(department: str = None, limit: int = 100) -> List[Dict]:
    try:
        client = _s3()
        prefix = f"uploads/{department}/" if department else "uploads/"
        resp   = client.list_objects_v2(Bucket=AWS_S3_BUCKET, Prefix=prefix, MaxKeys=limit)
        docs   = []
        for obj in resp.get('Contents', []):
            if obj['Key'].endswith('/'):
                continue
            try:
                meta = client.head_object(Bucket=AWS_S3_BUCKET, Key=obj['Key']).get('Metadata', {})
            except Exception:
                meta = {}
            docs.append({
                'key':           obj['Key'],
                'size':          obj['Size'],
                'last_modified': obj['LastModified'].isoformat(),
                'department':    meta.get('department', 'unknown'),
                'document_type': meta.get('document_type', 'unknown'),
                'url':           f"https://{AWS_S3_BUCKET}.s3.{AWS_REGION}.amazonaws.com/{obj['Key']}",
            })
        return docs
    except Exception as e:
        print(f"S3 list failed: {e}")
        return []

# ── Full document pipeline ────────────────────────────────────────────────────
def process_s3_document(s3_key: str) -> DocumentProcessingResult:
    print(f"🚀 Processing S3 document: {s3_key}")
    local_path        = download_from_s3(s3_key)
    original_filename = os.path.basename(s3_key)
    raw_text          = extract_text_from_file(local_path)
    doc_type          = classify_document(raw_text)
    department        = determine_department(doc_type, raw_text)
    summary           = create_summary(raw_text, doc_type)
    key_points        = extract_key_points(raw_text)
    action_items      = extract_action_items(raw_text)
    deadline          = extract_deadline(raw_text)
    priority          = determine_priority(raw_text)

    ts                 = datetime.now().strftime("%Y%m%d_%H%M%S")
    processed_filename = f"{department.value}_{ts}_{uuid.uuid4().hex[:8]}_{original_filename}"

    doc_metadata = {
        'original_filename':  original_filename,
        's3_key':             s3_key,
        'processed_date':     datetime.now().isoformat(),
        'document_type':      doc_type.value,
        'department':         department.value,
        'text_length':        len(raw_text),
        'priority':           priority,
        'has_deadline':       deadline is not None,
        'key_points_count':   len(key_points),
        'action_items_count': len(action_items),
    }

    s3_result = upload_to_s3(local_path, department.value, doc_type.value)
    os.unlink(local_path)

    return DocumentProcessingResult(
        file_path=local_path,
        original_filename=original_filename,
        processed_filename=processed_filename,
        document_type=doc_type,
        department=department,
        summary=summary,
        key_points=key_points,
        action_items=action_items,
        deadline=deadline,
        priority=priority,
        metadata=doc_metadata,
        raw_text=raw_text[:1000],
        processed_date=datetime.now().isoformat(),
        s3_key=s3_result['key'],
        s3_url=s3_result['url'],
    )

def batch_process_s3_documents(s3_keys: List[str]) -> Dict:
    results = {dept: [] for dept in Department}
    for key in s3_keys:
        try:
            r = process_s3_document(key)
            results[r.department].append(r)
        except Exception as e:
            print(f"Failed to process {key}: {e}")
    return results

def auto_fetch_and_process(department: str = None) -> Dict:
    try:
        client  = _s3()
        prefix  = "uploads/"
        resp    = client.list_objects_v2(Bucket=AWS_S3_BUCKET, Prefix=prefix, MaxKeys=50)
        keys    = [
            obj['Key'] for obj in resp.get('Contents', [])
            if not obj['Key'].endswith('/')
            and 'processed/' not in obj['Key']
            and (not department or f"/{department}/" in obj['Key'])
        ]
        if not keys:
            return {'message': 'No unprocessed documents found', 'processed': 0}

        results_by_dept = batch_process_s3_documents(keys[:10])

        for key in keys[:10]:
            try:
                archive_key = f"archive/{datetime.now().strftime('%Y/%m/%d')}/{os.path.basename(key)}"
                client.copy_object(Bucket=AWS_S3_BUCKET,
                                   CopySource={'Bucket': AWS_S3_BUCKET, 'Key': key},
                                   Key=archive_key)
                client.delete_object(Bucket=AWS_S3_BUCKET, Key=key)
            except Exception as e:
                print(f"Archive failed for {key}: {e}")

        total = sum(len(v) for v in results_by_dept.values())
        return {
            'message':        f'Processed {total} documents',
            'total_processed': total,
            'by_department':  {d.value: len(v) for d, v in results_by_dept.items() if v},
            'documents': [
                {'original_filename': r.original_filename, 'department': r.department.value,
                 'document_type': r.document_type.value, 'priority': r.priority, 's3_url': r.s3_url}
                for docs in results_by_dept.values() for r in docs
            ],
        }
    except Exception as e:
        return {'error': str(e), 'processed': 0}